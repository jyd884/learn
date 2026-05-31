import os

import openpyxl
from openpyxl.styles import Alignment

# 下面的方法一定要确保给定的 "工序记录明细查询表" 中每 30 行数据代表一个产品的生产记录, 否则就不能用

# 记录表中需要的字段
need_fields = ['批号', '缆芯外径(mm)', '护套外径(mm)', 'diff-输入(d)', '挤出内模(mm)', '挤出外模(mm)', 'diff-挤出(d)',
               '螺杆速度(rpm)-(挤塑主机速度)(转/分)', '螺杆电流(A)', '实际生产速度(m/min)', '初始行号', '来源文件名']


def get_project_root():
    return os.path.dirname(os.path.abspath(__file__))


# 读取给定工序记录明细查询表, 得到对应的所有记录
def getAllRecords(file_path):
    # 加载 Excel 文件
    projDir = get_project_root()
    wb = openpyxl.load_workbook(projDir + '/' + file_path)
    sheet = wb.active

    # allRecords 用于存放所有的记录, key 是 HT-xx, value 是对应的列表记录
    # 每一条记录都有类似于 ['批号', '缆芯外径(mm)', '护套外径(mm)', 'diff-输入(d)'...] 的字段
    allRecords = {}
    oneRecord = {}  # oneRecord 用于存放一条记录
    col = 23  # 只需要读取第 23 列(项目名称) 和第 24 列(项目记录结果)的值

    recordTotalNum = int((sheet.max_row - 1) / 30)
    for i in range(0, recordTotalNum):
        oneRecord['批号'] = sheet.cell(i * 30 + 2, 1).value
        devName = sheet.cell(i * 30 + 2, 14).value  # 设备名称
        oneRecord['初始行号'] = i * 30 + 2
        oneRecord['来源文件名'] = file_path
        # 给定的 excel 表中每 30 行数据为一个班次
        for j in range(1, 31):
            row = i * 30 + 1 + j
            projectName = sheet.cell(row, col).value
            # 将表格中字段的中文括号全转换为英文括号
            projectName = projectName.replace('（', '(')
            projectName = projectName.replace('）', ')')
            if projectName in need_fields:
                oneRecord[projectName] = sheet.cell(row, col + 1).value
        if devName not in allRecords.keys():
            allRecords[devName] = []
        allRecords[devName].append(oneRecord.copy())  # 将一个班次的信息保存到 allRecords 中
        oneRecord.clear()

    return allRecords


# 创建一个新的记录表
def createRecordExcel(file_path):
    # 如果记录表存在, 直接退出
    if os.path.exists(file_path):
        return
    # 创建一个新的工作簿
    wb = openpyxl.Workbook()
    # 选择默认的工作表
    sheet = wb.active
    # 指定每列的宽度
    column_widths = [20, 15, 15, 15, 15, 15, 15, 40, 15, 20, 10, 50]
    # 写入列名、设置列宽、设置居中对齐
    # need_fields = ['批号', '缆芯外径(mm)', '护套外径(mm)', 'diff-输入(d)', '挤出内模(mm)', '挤出外模(mm)', 'diff-挤出(d)',
    #                '螺杆速度(rpm)-(挤塑主机速度)(转/分)', '螺杆电流(A)', '实际生产速度(m/min)', '初始行号', '来源文件名']
    for col_num, column_name in enumerate(need_fields, start=1):
        sheet.cell(row=1, column=col_num, value=column_name)
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = column_widths[col_num - 1]
        for cell in sheet[
            openpyxl.utils.get_column_letter(col_num) + '1:' + openpyxl.utils.get_column_letter(col_num) + str(
                sheet.max_row)]:
            cell[0].alignment = Alignment(horizontal='center', vertical='center')
    # 保存工作簿
    wb.save(file_path)


# 判断记录中的字段是否有空值
def isNull(record):
    for filed in ['批号', '缆芯外径(mm)', '护套外径(mm)', '挤出内模(mm)', '挤出外模(mm)', '螺杆速度(rpm)-(挤塑主机速度)(转/分)',
                  '螺杆电流(A)', '实际生产速度(m/min)', '初始行号']:
        if filed not in record or record[filed] == '':
            return True
    return False


# 将所有记录写入对应的记录表中, 并去除字段为空的记录
def writeRecordToExcelAndRemoveNull(file_path, allRecords):
    for key in allRecords.keys():
        pos = file_path.find('/')
        fileName = file_path[:pos + 1] + key + '_' + file_path[pos + 1:]
        createRecordExcel(fileName)

        # 将 allRecords 中的记录放入一个二维列表中
        data = []
        for record in allRecords[key]:
            row_data = []
            if isNull(record):
                continue

            for filed in need_fields:
                if filed == 'diff-输入(d)':
                    row_data.append(float(record['护套外径(mm)']) - float(record['缆芯外径(mm)']))
                elif filed == 'diff-挤出(d)':
                    row_data.append(float(record['挤出外模(mm)']) - float(record['挤出内模(mm)']))
                else:
                    row_data.append(record[filed])
            data.append(row_data)

        # 打开 Excel 表
        wb = openpyxl.load_workbook(fileName)
        # 选择要操作的工作表
        sheet = wb.active
        # 将二维列表中的数据追加到 Excel 表中
        for row_data in data:
            sheet.append(row_data)
        # 保存修改后的 Excel 文件
        wb.save(fileName)


# 将所有记录写入对应的记录表中, 不去除字段为空的记录
def writeRecordToExcel(file_path, allRecords):
    for key in allRecords.keys():
        pos = file_path.find_last('/')
        fileName = file_path[:pos + 1] + key + '_' + file_path[pos + 1:]
        createRecordExcel(fileName)

        # 将 allRecords 中的记录放入一个二维列表中
        data = []
        for record in allRecords[key]:
            row_data = []
            for filed in need_fields:
                if filed == 'diff-输入(d)':
                    if '护套外径(mm)' not in record or '缆芯外径(mm)' not in record:
                        row_data.append('N/A')
                    else:
                        row_data.append(float(record['护套外径(mm)']) - float(record['缆芯外径(mm)']))
                elif filed == 'diff-挤出(d)':
                    if '挤出外模(mm)' not in record or '挤出内模(mm)' not in record:
                        row_data.append('N/A')
                    else:
                        row_data.append(float(record['挤出外模(mm)']) - float(record['挤出内模(mm)']))
                elif filed not in record:
                    row_data.append('N/A')
                else:
                    row_data.append(record[filed])
            data.append(row_data)

        # 打开 Excel 表
        wb = openpyxl.load_workbook(fileName)
        # 选择要操作的工作表
        sheet = wb.active
        # 将二维列表中的数据追加到 Excel 表中
        for row_data in data:
            sheet.append(row_data)
        # 保存修改后的 Excel 文件
        wb.save(fileName)


if __name__ == '__main__':
    files = ['original_data/HT_工序纪录明细查询(6月).xlsx', 'original_data/HT_工序纪录明细查询(7月).xlsx',
             'original_data/HT_工序纪录明细查询(8月).xlsx']
    # files = ['original_data/HT_工序纪录明细查询(6月).xlsx']

    for file in files:
        records = getAllRecords(file)
        writeRecordToExcelAndRemoveNull('bak/processed_data/final.xlsx', records)
        # writeRecordToExcel('processed_data/withNull/finalWithNull.xlsx', records)
